#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""金九银十·贵阳门店动作积分表(9月) → 看板数据JSON
用法: python3 jyhs_jifen_parse.py <积分表.xlsx> [out.json]
结构:
 每日存根: row1表头(0区域1军长2师长3渠道4门店分组5区域6军长7师长8渠道9门店分组10匹配词11门店 12..39日期 40合计), 数据 row4+
 每日公示: 同前6列 + 7..18 = 12个动作 19合计
"""
import openpyxl, json, sys, warnings
from collections import defaultdict
from datetime import datetime, timedelta
warnings.filterwarnings('ignore')
BASE = datetime(1899, 12, 30)

ACTIONS = ['晨会','夕会','每日一读','产品演练','上门量尺','KDS预约','新增客资','门店物料落地','捷报','会议签到','周末活动落地','渠道活动落地']

def num(v):
    return v if isinstance(v, (int, float)) else 0

def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['每日存根']
    rows = list(ws.iter_rows(values_only=True))
    header = rows[1]
    # 日期列
    date_cols = []
    for j, v in enumerate(header):
        if isinstance(v, (int, float)) and v > 40000:
            date_cols.append(j)
    labels = [(BASE + timedelta(days=header[j])).strftime('%m/%d') for j in date_cols]
    n = len(date_cols)
    stores = []
    for r in rows[4:]:
        if not r or not isinstance(r[11], str) or not str(r[11]).strip():
            continue
        store = str(r[11]).strip()
        if '合计' in store:
            continue
        daily = [num(r[j]) for j in date_cols]
        if sum(daily) == 0 and not any(r[k] for k in (0,)):
            pass
        stores.append({
            'region': r[0], 'jun': r[1], 'shi': r[2], 'ch': r[3], 'group': r[4],
            'store': store, 'daily': daily, 'total': sum(daily),
        })

    # 去除标题行占位/空行（门店名为空已跳过）
    n_days = n
    daily_total = [sum(s['daily'][i] for s in stores) for i in range(n_days)]
    reg_daily = defaultdict(lambda: [0]*n_days)
    reg_cnt = defaultdict(int)
    for s in stores:
        reg_cnt[s['region'] or '未分类'] += 1
        for i, d in enumerate(s['daily']):
            reg_daily[s['region'] or '未分类'][i] += d

    # 组别(S/A/B)
    groups = defaultdict(list)
    for s in stores:
        groups[s['group'] or '未分组'].append(s)
    group_tot = {}
    for g, arr in groups.items():
        group_tot[g] = {'count': len(arr), 'total': sum(x['total'] for x in arr),
                        'avg': (sum(x['total'] for x in arr)/len(arr)) if arr else 0,
                        'stores': sorted(arr, key=lambda x: -x['total'])}

    # 今日动作（每日公示）
    today = []
    try:
        wst = wb['每日公示']
        trs = list(wst.iter_rows(values_only=True))
        for r in trs[4:]:
            if not r or not isinstance(r[6], str) or not str(r[6]).strip():
                continue
            store = str(r[6]).strip()
            if '合计' in store:
                continue
            acts = {}
            for i, name in enumerate(ACTIONS):
                acts[name] = num(r[7+i])
            today.append({'region': r[0], 'store': store, 'acts': acts, 'total': num(r[19])})
    except Exception as e:
        print('每日公示解析失败:', e)

    act_done = defaultdict(int)
    for t in today:
        for a in ACTIONS:
            if t['acts'].get(a, 0) > 0:
                act_done[a] += 1

    data = {
        'kind': 'jyhs_points',
        'labels': labels, 'n_days': n_days,
        'stores': stores,
        'daily_total': daily_total,
        'reg_daily': dict(reg_daily),
        'reg_cnt': dict(reg_cnt),
        'reg_total': {k: sum(v) for k, v in reg_daily.items()},
        'groups': {g: {'count': v['count'], 'total': v['total'], 'avg': v['avg'],
                       'stores': [{'store': x['store'], 'total': x['total'], 'group': x['group']} for x in v['stores']]}
                   for g, v in group_tot.items()},
        'today': today,
        'act_done': dict(act_done),
        'act_names': ACTIONS,
        'total_stores': len(stores),
        'total_score': sum(s['total'] for s in stores),
        'updated': datetime.now().strftime('%Y-%m-%d %H:%M'),
    }
    return data

if __name__ == '__main__':
    src = sys.argv[1]; out = sys.argv[2] if len(sys.argv) > 2 else 'jyhs_jifen.json'
    d = parse(src)
    json.dump(d, open(out, 'w'), ensure_ascii=False)
    print(f"✅ 9月积分解析: 门店{d['total_stores']}家 / {d['n_days']}天 / 总积分{d['total_score']} / 组别{list(d['groups'].keys())}")
    print(f"   有分数天数: {sum(1 for x in d['daily_total'] if x>0)} / 每日: {d['daily_total'][:8]}")
    print(f"   输出: {out}")
