#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
金九银十贵阳老板电器 · 项目数据回传表 → 业绩看板数据JSON
用法: python3 jyhs_parse.py <数据回传.xlsx> [输出.json]
表结构(Sheet1):
  col1区域 col2军长 col3师长 col4渠道 col5门店
  col6 9月目标 col7 9月完成 col8 完成度
  col9 一阶段目标(40%) col10 一阶段完成 col11 一阶段完成率
  col12 日目标 col13 日完成 col14 日完成率
"""
import openpyxl, json, sys, warnings
from collections import defaultdict
from datetime import datetime, timedelta
warnings.filterwarnings('ignore')

BASE = datetime(1899, 12, 30)

def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0

def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Sheet1']

    def g(r, c):
        return ws.cell(r, c).value

    # ---- 顶部汇总行（行3=张美娜市区, 行4=合计赖元章）----
    overall = None
    for r in (3, 4):
        if isinstance(g(r, 6), (int, float)) and g(r, 5):
            pass
    # 行4 合计
    t_all, d_all = num(g(4, 6)), num(g(4, 7))
    overall = {
        'target': t_all, 'done': d_all,
        'rate': d_all / t_all if t_all else 0,
        'diff': max(0, t_all - d_all),
        'teacher_target': num(g(3, 6)), 'teacher_done': num(g(3, 7)),
    }
    # 起止时间
    sd, ed = g(3, 2), g(3, 3)
    if isinstance(sd, (int, float)) and isinstance(ed, (int, float)):
        s = BASE + timedelta(days=int(sd)); e = BASE + timedelta(days=int(ed))
        today = datetime.now()
        total_days = (e - s).days or 1
        elapsed = max(0, min((today - s).days, total_days))
        overall['start'] = s.strftime('%m/%d'); overall['end'] = e.strftime('%m/%d')
        overall['start_full'] = s.strftime('%Y-%m-%d'); overall['end_full'] = e.strftime('%Y-%m-%d')
        overall['time_pct'] = round(elapsed / total_days * 100, 1)
        overall['remain_days'] = max(0, (e - today).days)

    # ---- 门店明细 ----
    stores = []
    cur = {'region': None, 'jun': None, 'shi': None, 'ch': None}
    for r in range(9, ws.max_row + 1):
        store = g(r, 5)
        if not (isinstance(store, str) and store.strip()):
            continue
        if g(r, 6) is None and g(r, 7) is None:
            continue
        for idx, key in ((1, 'region'), (2, 'jun'), (3, 'shi'), (4, 'ch')):
            v = g(r, idx)
            if isinstance(v, str) and v.strip():
                cur[key] = v.strip()
        t9, d9 = num(g(r, 6)), num(g(r, 7))
        pt, pd = num(g(r, 9)), num(g(r, 10))
        dt, dd = num(g(r, 12)), num(g(r, 13))
        if t9 == 0 and d9 == 0 and pt == 0:
            continue
        stores.append({
            'region': cur['region'], 'jun': cur['jun'], 'shi': cur['shi'], 'ch': cur['ch'],
            'store': store.strip(),
            't9': t9, 'd9': d9, 'r9': (d9 / t9 if t9 else 0),
            'pt': pt, 'pd': pd, 'pr': (pd / pt if pt else 0),
            'dt': dt, 'dd': dd, 'dr': (dd / dt if dt else 0),
        })

    def agg(key):
        tot = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])  # t9,d9,pt,pd
        for s in stores:
            k = s[key] or '未分类'
            tot[k][0] += s['t9']; tot[k][1] += s['d9']
            tot[k][2] += s['pt']; tot[k][3] += s['pd']
        out = []
        for k, (t, d, pt, pd) in tot.items():
            out.append({'name': k, 't9': t, 'd9': d, 'r9': d / t if t else 0,
                        'pt': pt, 'pd': pd, 'pr': pd / pt if pt else 0})
        out.sort(key=lambda x: -x['t9'])
        return out

    detail_total = sum(s['t9'] for s in stores)
    detail_done = sum(s['d9'] for s in stores)

    data = {
        'kind': 'jyhs_perf',
        'overall': overall,
        'detail_total': detail_total,
        'detail_done': detail_done,
        'by_jun': agg('jun'),
        'by_channel': agg('ch'),
        'by_region': agg('region'),
        'stores': stores,
        'updated': datetime.now().strftime('%Y-%m-%d %H:%M'),
    }
    return data

if __name__ == '__main__':
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else 'jyhs_perf.json'
    data = parse(src)
    with open(out, 'w') as f:
        json.dump(data, f, ensure_ascii=False)
    o = data['overall']
    print(f"✅ 金九银十业绩解析: 门店{len(data['stores'])}家 / 总目标{round(o['target']/10000,1)}万 / "
          f"完成{round(o['done']/10000,1)}万 / 达成率{round(o['rate']*100,2)}% / 时间进度{o.get('time_pct',0)}%")
    print(f"   输出: {out}")
